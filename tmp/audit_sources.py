from pathlib import Path
import sys, re
import xlrd
from openpyxl import load_workbook
base=Path('data_sources')
files=[]
for p in sorted(base.iterdir()):
    if p.suffix.lower() in ('.xls','.xlsx','.ods') and not p.name.startswith('HE_'):
        files.append(p)
for p in files:
    print('\n###',p.name)
    try:
      if p.suffix.lower()=='.xls':
        b=xlrd.open_workbook(str(p))
        for s in b.sheets():
          print('SHEET',repr(s.name),s.nrows,s.ncols)
          for r in range(min(8,s.nrows)):
            vals=[str(s.cell_value(r,c)).replace('\n',' ')[:80] for c in range(min(s.ncols,12))]
            if any(v not in ('','None') for v in vals): print(' ',r+1,vals)
      elif p.suffix.lower()=='.xlsx':
        b=load_workbook(p,read_only=True,data_only=True)
        for s in b.worksheets:
          print('SHEET',repr(s.title),s.max_row,s.max_column)
          for r,row in enumerate(s.iter_rows(min_row=1,max_row=min(8,s.max_row),values_only=True),1):
            vals=[str(x).replace('\n',' ')[:80] if x is not None else '' for x in row[:12]]
            if any(vals): print(' ',r,vals)
      else:
        import pandas as pd
        x=pd.ExcelFile(p,engine='odf')
        for sn in x.sheet_names:
          df=pd.read_excel(p,sheet_name=sn,header=None,engine='odf')
          print('SHEET',repr(sn),df.shape[0],df.shape[1])
          for r in range(min(8,len(df))):
            vals=[str(x).replace('\n',' ')[:80] if str(x)!='nan' else '' for x in df.iloc[r,:12]]
            if any(vals): print(' ',r+1,vals)
    except Exception as e:
      print('ERROR',type(e).__name__,e)
